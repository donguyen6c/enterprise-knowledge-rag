import json
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.core.files.base import ContentFile
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from docx import Document as DocxDocument

from accounts.models import User
from documents.models import (
    Document,
    DocumentCategory,
    DocumentPermission,
)
from documents.services.media_paths import normalize_document_file_path
from documents.services.processor import process_and_index_document
from organizations.models import Department, Organization


DEMO_DATA_DIR = settings.BASE_DIR.parent / "sample-data" / "kfc-demo"
MANIFEST_PATH = DEMO_DATA_DIR / "corpus.json"
DEFAULT_DEMO_PASSWORD = "KfcDemo@2026"


def load_manifest() -> dict:
    try:
        return json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise CommandError(f"Không tìm thấy manifest: {MANIFEST_PATH}") from exc
    except json.JSONDecodeError as exc:
        raise CommandError(f"Manifest KFC demo không hợp lệ: {exc}") from exc


def build_docx(document_data: dict) -> bytes:
    output = BytesIO()
    document = DocxDocument()
    document.core_properties.title = document_data["title"]
    document.add_heading(document_data["title"], level=0)

    notice = document_data.get("notice", "").strip()

    if notice:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(notice)
        run.bold = True

    for section in document_data.get("sections", []):
        document.add_heading(section["heading"], level=1)

        for text in section.get("paragraphs", []):
            document.add_paragraph(text)

    sources = document_data.get("sources", [])

    if sources:
        document.add_heading("Nguồn tham khảo", level=1)

        for source in sources:
            document.add_paragraph(
                f"{source['title']}\n"
                f"{source['url']}\n"
                f"Ngày truy cập: {source['accessed']}"
            )

    document.save(output)
    return output.getvalue()


def read_document_bytes(document_data: dict) -> bytes:
    source_file = document_data.get("source_file")

    if not source_file:
        return build_docx(document_data)

    data_root = DEMO_DATA_DIR.resolve()
    file_path = (data_root / source_file).resolve()

    if not file_path.is_relative_to(data_root):
        raise CommandError(f"Đường dẫn source_file không hợp lệ: {source_file}")

    if not file_path.is_file():
        raise CommandError(f"Không tìm thấy source_file: {file_path}")

    return file_path.read_bytes()


class Command(BaseCommand):
    help = "Tạo hoặc xóa tổ chức KFC Việt Nam dùng cho demo RBAC."

    def add_arguments(self, parser):
        parser.add_argument(
            "--password",
            default=DEFAULT_DEMO_PASSWORD,
            help="Mật khẩu dùng chung cho ba tài khoản KFC demo.",
        )
        parser.add_argument(
            "--remove",
            action="store_true",
            help="Xóa riêng tổ chức, user, tài liệu và media KFC demo.",
        )
    def handle(self, *args, **options):
        manifest = load_manifest()
        organization_name = manifest["organization"]["name"]

        if options["remove"]:
            self.remove_demo(organization_name)
            return

        password = options["password"]

        if len(password) < 10:
            raise CommandError("Mật khẩu demo phải có ít nhất 10 ký tự.")

        organization = self.upsert_organization(manifest["organization"])
        departments = self.upsert_departments(
            organization,
            manifest["departments"],
        )
        users = self.upsert_users(
            organization,
            departments,
            manifest["users"],
            password,
        )
        processed_documents = self.upsert_documents(
            organization,
            departments,
            users,
            manifest["documents"],
        )

        self.stdout.write("")
        self.stdout.write(
            self.style.SUCCESS(
                f"KFC demo sẵn sàng: organization_id={organization.id}, "
                f"documents={processed_documents}."
            )
        )
        self.stdout.write("Tài khoản demo:")

        for user_data in manifest["users"]:
            self.stdout.write(f"- {user_data['email']}")

        self.stdout.write(f"Mật khẩu dùng chung: {password}")

    def upsert_organization(self, data: dict) -> Organization:
        organization = Organization.objects.filter(name=data["name"]).first()

        if organization is None:
            organization = Organization(name=data["name"])

        organization.description = data.get("description", "")
        organization.is_active = True
        organization.full_clean()
        organization.save()
        return organization

    def upsert_departments(
        self,
        organization: Organization,
        department_data: list[dict],
    ) -> dict[str, Department]:
        departments = {}

        for data in department_data:
            department, _ = Department.objects.get_or_create(
                organization=organization,
                name=data["name"],
            )
            department.description = data.get("description", "")
            department.is_active = True
            department.full_clean()
            department.save()
            departments[data["key"]] = department

        return departments

    def upsert_users(
        self,
        organization: Organization,
        departments: dict[str, Department],
        user_data: list[dict],
        password: str,
    ) -> dict[str, User]:
        users = {}

        for data in user_data:
            user = User.objects.filter(email=data["email"]).first()

            if user is None:
                user = User(email=data["email"])

            user.username = data["username"]
            user.first_name = data.get("first_name", "")
            user.last_name = data.get("last_name", "")
            user.organization = organization
            user.department = departments[data["department_key"]]
            user.role = data["role"]
            user.is_active = True
            user.set_password(password)
            user.full_clean()
            user.save()
            users[data["key"]] = user

        return users

    def upsert_documents(
        self,
        organization: Organization,
        departments: dict[str, Department],
        users: dict[str, User],
        document_data: list[dict],
    ) -> int:
        processed_count = 0

        for data in document_data:
            category, _ = DocumentCategory.objects.get_or_create(
                organization=organization,
                name=data["category"],
            )
            document = Document.objects.filter(
                organization=organization,
                title=data["title"],
            ).first()

            if document is None:
                document = Document(
                    organization=organization,
                    title=data["title"],
                )

            file_bytes = read_document_bytes(data)
            original_filename = data["original_filename"]
            document.category = category
            document.uploaded_by = users[data["uploader_key"]]
            document.description = data.get("description", "")
            document.visibility = data["visibility"]
            document.status = Document.Status.UPLOADED
            document.original_filename = original_filename
            document.file_size = len(file_bytes)
            document.error_message = ""
            document.is_active = True
            document.file.save(
                original_filename,
                ContentFile(file_bytes),
                save=False,
            )
            document.full_clean()
            document.save()

            normalization = normalize_document_file_path(
                document,
                apply=True,
                replace_target=True,
            )

            if normalization.changed and not normalization.applied:
                raise CommandError(
                    f"Không chuẩn hóa được media của document #{document.id}."
                )

            document.permissions.all().delete()
            permission_data = data.get("permission")

            if permission_data:
                permission = DocumentPermission(
                    document=document,
                    can_view=True,
                    can_download=True,
                )

                if permission_data.get("department_key"):
                    permission.department = departments[
                        permission_data["department_key"]
                    ]
                else:
                    permission.role = permission_data["role"]

                permission.full_clean()
                permission.save()

            chunk_count, embedding_count = process_and_index_document(document)
            processed_count += 1
            self.stdout.write(
                self.style.SUCCESS(
                    f"READY #{document.id}: {document.title} "
                    f"({chunk_count} chunks, {embedding_count} embeddings)"
                )
            )

        return processed_count

    def remove_demo(self, organization_name: str) -> None:
        organizations = list(
            Organization.objects.filter(name=organization_name)
        )

        if not organizations:
            self.stdout.write(self.style.WARNING("KFC demo không tồn tại."))
            return

        file_references = []

        for organization in organizations:
            for document in organization.documents.exclude(file=""):
                file_references.append(
                    (document.file.storage, document.file.name)
                )

        with transaction.atomic():
            for organization in organizations:
                organization.documents.all().delete()
                organization.users.all().delete()
                organization.delete()

        for storage, file_name in file_references:
            storage.delete(file_name)

        self.stdout.write(
            self.style.SUCCESS(
                f"Đã xóa {len(organizations)} tổ chức KFC demo và "
                f"{len(file_references)} file media."
            )
        )
